import io

from docx import Document

from app.ai.base import LLMProvider
from app.ai.schemas import (
    ResumeCertificationLLM,
    ResumeEducationLLM,
    ResumeExperienceLLM,
    ResumeLLMExtraction,
)
from app.services import resume_service

RESUME_TEXT_LINES = [
    "Priya Raman",
    "Education",
    "B.Tech in Computer Science, National Institute of Technology Trichy, 2023",
    "Experience",
    "Machine Learning Engineer, Zoho Corporation, 2023-07 to present",
    "Built YOLO-based defect detection with PyTorch and OpenCV.",
    "Certifications",
    "JLPT N3",
]


def _docx_bytes(lines: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _FakeResumeLLM(LLMProvider):
    is_llm = True

    def __init__(self, extraction):
        self._extraction = extraction

    def extract_requirements(self, job_title, description, language):
        raise NotImplementedError

    def explain_match(self, job_title, company_name, requirement_breakdown, match_strength):
        raise NotImplementedError

    def extract_resume(self, resume_text):
        return self._extraction


def _llm_result(**overrides) -> ResumeLLMExtraction:
    defaults = dict(
        education=[
            ResumeEducationLLM(
                degree="B.Tech", university="National Institute of Technology Trichy",
                field_of_study="Computer Science", graduation_year=2023, is_current_student=False,
            )
        ],
        experience=[
            ResumeExperienceLLM(
                company_name="Zoho Corporation", role="Machine Learning Engineer", is_internship=False,
                start_date="2023-07", end_date=None, description="Built defect detection models.",
            )
        ],
        projects=[],
        certifications=[ResumeCertificationLLM(name="JLPT N3", issuer=None, year=None)],
        skills=["pytorch", "OpenCV", "YOLO"],
        jlpt_level="N3",
        total_years_of_experience=2.2,
    )
    defaults.update(overrides)
    return ResumeLLMExtraction(**defaults)


def test_docx_text_includes_tables():
    content = _docx_bytes(["Skills"], table_rows=[["Languages", "Python, Go"], ["Frameworks", "FastAPI"]])
    text = resume_service.extract_text("cv.docx", content)
    assert "Languages | Python, Go" in text
    assert "FastAPI" in text


def test_llm_extraction_is_used_and_skills_are_canonicalized(monkeypatch):
    monkeypatch.setattr(resume_service, "get_resume_llm_provider", lambda: _FakeResumeLLM(_llm_result()))

    result = resume_service.parse_resume("cv.docx", _docx_bytes(RESUME_TEXT_LINES))

    assert result.extraction_method == "ai"
    assert [e.company_name for e in result.experience] == ["Zoho Corporation"]
    assert [s.name for s in result.skills][:3] == ["PyTorch", "OpenCV", "YOLO"]
    assert result.jlpt_level.value == "N3"
    assert result.years_of_experience == 2.2


def test_items_not_in_the_resume_are_dropped(monkeypatch):
    hallucinated = _llm_result(
        experience=[
            ResumeExperienceLLM(
                company_name="Google", role="Senior Engineer", is_internship=False,
                start_date="2019-01", end_date="2022-12", description=None,
            )
        ],
        skills=["PyTorch", "Kubernetes"],
        jlpt_level="N1",
    )
    monkeypatch.setattr(resume_service, "get_resume_llm_provider", lambda: _FakeResumeLLM(hallucinated))

    result = resume_service.parse_resume("cv.docx", _docx_bytes(RESUME_TEXT_LINES))

    assert result.experience == []
    assert "Kubernetes" not in [s.name for s in result.skills]
    assert result.jlpt_level is None
    assert any("left out 1 item" in w for w in result.warnings)


def test_falls_back_to_heuristics_when_llm_fails(monkeypatch):
    monkeypatch.setattr(resume_service, "get_resume_llm_provider", lambda: _FakeResumeLLM(None))

    result = resume_service.parse_resume("cv.docx", _docx_bytes(RESUME_TEXT_LINES))

    assert result.extraction_method == "rules"
    assert {"PyTorch", "OpenCV", "YOLO"} <= {s.name for s in result.skills}
    assert result.jlpt_level.value == "N3"


def test_years_of_experience_derived_from_roles():
    from datetime import date

    from app.schemas.profile import ExperienceIn
    from app.services.profile_service import derive_years_of_experience

    roles = [
        ExperienceIn(company_name="A", role="Engineer", start_date="2020-01", end_date="2022-01"),
        # Overlaps the first role by a year: counted once.
        ExperienceIn(company_name="B", role="Engineer", start_date="2021-01", end_date="2023-01"),
        ExperienceIn(company_name="C", role="Intern", is_internship=True, start_date="2018-01", end_date="2019-01"),
    ]
    assert derive_years_of_experience(roles, today=date(2026, 1, 1)) == 3.0
    assert derive_years_of_experience([ExperienceIn(company_name="A", role="x")]) is None
